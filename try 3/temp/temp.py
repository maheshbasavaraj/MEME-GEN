from docx import Document

# Create the Word document for AIML lesson plan
doc = Document()
doc.add_heading("Lesson Plan: Artificial Intelligence and Machine Learning (Theory)", level=1)

# Define the table headers
headers = [
    "Module No", "Lecture No.", "Session Topics", "RBT Levels",
    "Activities Planned", "Course Outcome Mapping", "Mode of Delivery",
    "Planned Date", "Actual Date"
]

# Add the table
table = doc.add_table(rows=1, cols=len(headers))
table.style = "Table Grid"
hdr_cells = table.rows[0].cells
for i, header in enumerate(headers):
    hdr_cells[i].text = header

# Define detailed sessions for each AIML module (8 sessions each)
aiml_lessons = {
    1: [
        ("What is AI? Definitions and History", "Understand (L2)", "Chalk and Talk, PPT", "CO1"),
        ("Current Trends: Generative AI & Agentic Systems", "Understand (L2)", "Group Discussion", "CO1"),
        ("Intelligent Agents: Types and Characteristics", "Understand (L2)", "Think-Pair-Share", "CO1"),
        ("Rationality and Environments", "Understand (L2)", "Case Study Analysis", "CO1"),
        ("Problem-Solving Frameworks", "Understand (L2)", "Chalk and Talk, PPT", "CO2"),
        ("State-Space Search: Breadth-First Search (BFS)", "Apply (L3)", "Collaborative Problem Solving", "CO2"),
        ("State-Space Search: Depth-First Search (DFS)", "Apply (L3)", "Collaborative Problem Solving", "CO2"),
        ("Recap and Q&A", "Analyze (L4)", "Discussion", "CO2"),
    ],
    2: [
        ("Intro to ML: Supervised, Unsupervised, Reinforcement", "Understand (L2)", "Chalk and Talk, PPT", "CO2"),
        ("ML Process Overview", "Understand (L2)", "Chalk and Talk, PPT", "CO2"),
        ("Data Preprocessing Techniques", "Apply (L3)", "Problem-Solving", "CO2"),
        ("Feature Engineering Techniques", "Apply (L3)", "Hands-on Demo", "CO2"),
        ("Model Evaluation: Accuracy & Precision", "Analyze (L4)", "Quiz", "CO2"),
        ("Model Evaluation: Recall & F1 Score", "Analyze (L4)", "Quiz", "CO2"),
        ("Data Types and Descriptive Statistics", "Understand (L2)", "Chalk and Talk, PPT", "CO2"),
        ("Introduction to Data Visualization", "Apply (L3)", "Demonstration", "CO2"),
    ],
    3: [
        ("Linear Regression: Concept and Theory", "Understand (L2)", "Chalk and Talk, PPT", "CO3"),
        ("Linear Regression: Implementation Demo", "Apply (L3)", "Hands-on Demo", "CO3"),
        ("Logistic Regression: Concept and Use-Cases", "Understand (L2)", "Chalk and Talk, PPT", "CO3"),
        ("Decision Trees: Theory and Splitting", "Understand (L2)", "Chalk and Talk, PPT", "CO3"),
        ("Decision Trees: Implementation Demo", "Apply (L3)", "Hands-on Demo", "CO3"),
        ("K-Nearest Neighbors: Concept & Demo", "Apply (L3)", "Collaborative Problem Solving", "CO3"),
        ("K-Means Clustering: Theory & Demo", "Apply (L3)", "Hands-on Demo", "CO3"),
        ("Principal Component Analysis (PCA): Overview", "Understand (L2)", "Chalk and Talk, PPT", "CO3"),
    ],
    4: [
        ("Introduction to Neural Networks & Perceptron", "Understand (L2)", "Chalk and Talk, PPT", "CO4"),
        ("MLP Architecture and Layers", "Understand (L2)", "Chalk and Talk, PPT", "CO4"),
        ("Activation Functions", "Understand (L2)", "Group Activity", "CO4"),
        ("Backpropagation: Theory", "Understand (L2)", "Chalk and Talk, PPT", "CO4"),
        ("Backpropagation: Worked Example", "Apply (L3)", "Collaborative Problem Solving", "CO4"),
        ("Training Process: Loss, Epochs, Batch Size", "Understand (L2)", "Demonstration", "CO4"),
        ("Overfitting and Regularization Techniques", "Analyze (L4)", "Case Study", "CO4"),
        ("Recap and Neural Network Discussion", "Analyze (L4)", "Discussion", "CO4"),
    ],
    5: [
        ("Generative Adversarial Networks (GANs): Theory", "Understand (L2)", "Chalk and Talk, PPT", "CO5"),
        ("GANs: Applications and Examples", "Understand (L2)", "Case Study", "CO5"),
        ("Variational Autoencoders (VAEs): Concepts", "Understand (L2)", "Chalk and Talk, PPT", "CO5"),
        ("Diffusion Models: Overview", "Understand (L2)", "Chalk and Talk, PPT", "CO5"),
        ("AGI: Definitions and Societal Impact", "Analyze (L4)", "Case Study", "CO5"),
        ("ASI: Definitions and Future Considerations", "Analyze (L4)", "Case Study", "CO5"),
        ("Ethical Implications of AI: Debate", "Evaluate (L5)", "Structured Debate", "CO5"),
        ("Deployment Strategies for AI Systems", "Apply (L3)", "Discussion", "CO5"),
    ],
}

# Populate the table
for module, sessions in aiml_lessons.items():
    for i, (topic, rbt, activity, co) in enumerate(sessions, start=1):
        row = table.add_row().cells
        row[0].text = str(module)
        row[1].text = str(i)
        row[2].text = topic
        row[3].text = rbt
        row[4].text = activity
        row[5].text = co
        row[6].text = "Chalk and Talk, PPT"
        row[7].text = ""  # Planned Date
        row[8].text = ""  # Actual Date

# Save document
aiml_lesson_plan_path = "AIML_Lesson_Plan_classes.docx"
doc.save(aiml_lesson_plan_path)

aiml_lesson_plan_path
